import os
import json
from pathlib import Path

try:
    from docx import Document
except ImportError:
    Document = None

class PrismaManager:
    def __init__(self, output_dir: str = None):
        if output_dir is None:
            # Default: data/ directory relative to project root
            data_dir = Path(__file__).resolve().parent.parent.parent / "data"
        else:
            out_path = Path(output_dir)
            # If output_dir looks like project root or relatorio, find data/ accordingly
            data_candidate = out_path / "data"
            if data_candidate.exists() or not (out_path.parent / "data").exists():
                data_dir = data_candidate
            else:
                data_dir = out_path.parent / "data"
        self.state_file = data_dir / "prisma_state.json"
        self.state_file.parent.mkdir(parents=True, exist_ok=True)
        self.state = self._load_state()

    def _load_state(self):
        if self.state_file.exists():
            with open(self.state_file, "r", encoding="utf-8") as f:
                return json.load(f)
        return {
            "identification": {"databases": [], "n_found": 0, "n_other": 0, "n_duplicates": 0},
            "screening": {"n_screened": 0, "n_excluded": 0},
            "retrieval": {"n_sought": 0, "n_not_retrieved": 0},
            "eligibility": {"n_assessed": 0, "n_excluded": 0, "exclusion_reasons": {}},
            "included": {"n_included": 0}
        }

    def _save_state(self):
        with open(self.state_file, "w", encoding="utf-8") as f:
            json.dump(self.state, f, indent=2)

    def update_identification(self, databases: list, n_found: int, n_other: int, n_duplicates: int):
        self.state["identification"] = {
            "databases": databases,
            "n_found": n_found,
            "n_other": n_other,
            "n_duplicates": n_duplicates
        }
        self._save_state()

    def update_screening(self, n_screened: int, n_excluded: int):
        self.state["screening"] = {
            "n_screened": n_screened,
            "n_excluded": n_excluded
        }
        self._save_state()
        
    def update_retrieval(self, n_sought: int, n_not_retrieved: int):
        self.state["retrieval"] = {
            "n_sought": n_sought,
            "n_not_retrieved": n_not_retrieved
        }
        self._save_state()

    def update_included(self, n_included: int, n_excluded_fulltext: int, exclusion_reasons: dict):
        self.state["eligibility"] = {
            "n_assessed": n_included + n_excluded_fulltext,
            "n_excluded": n_excluded_fulltext,
            "exclusion_reasons": exclusion_reasons
        }
        self.state["included"] = {"n_included": n_included}
        self._save_state()

    def export(self, out_dir: str):
        out_path = Path(out_dir)
        out_path.mkdir(parents=True, exist_ok=True)
        
        md_path = out_path / "PRISMA_2020.md"
        docx_path = out_path / "PRISMA_2020.docx"
        
        # 1. Gerar Markdown
        with open(md_path, "w", encoding="utf-8") as f:
            f.write("# Fluxograma PRISMA 2020 (Automático)\n\n")
            
            f.write("## 1. Identificação de estudos via bases de dados e registros\n")
            f.write(f"- Bases pesquisadas: {', '.join(self.state['identification']['databases'])}\n")
            f.write(f"- Registros identificados em bases de dados (n = {self.state['identification']['n_found']})\n")
            f.write(f"- Registros identificados por outros métodos (n = {self.state['identification']['n_other']})\n")
            f.write(f"- Registros removidos antes da triagem (Duplicatas) (n = {self.state['identification']['n_duplicates']})\n\n")
            
            f.write("## 2. Triagem (Screening)\n")
            f.write(f"- Registros triados (n = {self.state['screening']['n_screened']})\n")
            f.write(f"- Registros excluídos na triagem inicial (n = {self.state['screening']['n_excluded']})\n")
            f.write(f"- Relatórios buscados para recuperação (n = {self.state['retrieval']['n_sought']})\n")
            f.write(f"- Relatórios não recuperados/baixados (n = {self.state['retrieval']['n_not_retrieved']})\n\n")
            
            f.write("## 3. Elegibilidade e Inclusão\n")
            f.write(f"- Relatórios avaliados para elegibilidade em texto completo (n = {self.state['eligibility']['n_assessed']})\n")
            f.write(f"- Relatórios excluídos (n = {self.state['eligibility']['n_excluded']})\n")
            for reason, count in self.state['eligibility']['exclusion_reasons'].items():
                f.write(f"    - Razão: {reason} (n = {count})\n")
            f.write(f"- **Estudos incluídos na revisão (n = {self.state['included']['n_included']})**\n")

        # 2. Gerar Docx
        if Document:
            doc = Document()
            doc.add_heading('Fluxograma PRISMA 2020', 0)
            
            doc.add_heading('1. Identificação', level=1)
            doc.add_paragraph(f"Bases pesquisadas: {', '.join(self.state['identification']['databases'])}")
            doc.add_paragraph(f"Registros identificados (n = {self.state['identification']['n_found']})", style='List Bullet')
            doc.add_paragraph(f"Registros via outros métodos (n = {self.state['identification']['n_other']})", style='List Bullet')
            doc.add_paragraph(f"Duplicatas removidas (n = {self.state['identification']['n_duplicates']})", style='List Bullet')
            
            doc.add_heading('2. Triagem', level=1)
            doc.add_paragraph(f"Registros triados (n = {self.state['screening']['n_screened']})", style='List Bullet')
            doc.add_paragraph(f"Registros excluídos (n = {self.state['screening']['n_excluded']})", style='List Bullet')
            doc.add_paragraph(f"Buscados para recuperação (n = {self.state['retrieval']['n_sought']})", style='List Bullet')
            doc.add_paragraph(f"Não recuperados/baixados (n = {self.state['retrieval']['n_not_retrieved']})", style='List Bullet')
            
            doc.add_heading('3. Elegibilidade e Inclusão', level=1)
            doc.add_paragraph(f"Avaliados em texto completo (n = {self.state['eligibility']['n_assessed']})", style='List Bullet')
            doc.add_paragraph(f"Excluídos (n = {self.state['eligibility']['n_excluded']})", style='List Bullet')
            for reason, count in self.state['eligibility']['exclusion_reasons'].items():
                doc.add_paragraph(f"Razão: {reason} (n = {count})", style='List Bullet 2')
                
            p = doc.add_paragraph()
            p.add_run(f"Estudos incluídos na revisão (n = {self.state['included']['n_included']})").bold = True
            
            try:
                doc.save(docx_path)
            except Exception as e:
                print(f"Erro ao salvar DOCX: {e}")
